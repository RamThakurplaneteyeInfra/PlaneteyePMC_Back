"""Physical progress chart — missing reporting-month presentation tests."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import SimpleTestCase, TestCase
from openpyxl import load_workbook
from pypdf import PdfReader

from construction_progress.models.construction_progress import ConstructionProgress
from monthly_scope.models import MonthlyScopeWork, ScopeCategory, ScopeSubCategory
from mpr.services.excel_renderer import render_excel
from mpr.services.mpr_service import MPRService
from mpr.services.mpr_validation import validate_and_normalize_snapshot
from mpr.services.period import parse_mpr_month
from mpr.services.pdf_renderer import render_pdf
from mpr.tests_client_format import _minimal_snapshot
from projects.models import Project

User = get_user_model()


def _pdf_text(pdf_bytes: bytes) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf_bytes)).pages)


class PhysicalChartPresentationUnitTests(SimpleTestCase):
    def test_only_may_cp_july_report_no_july_point(self):
        snap = _minimal_snapshot()
        snap["physical_progress"]["monthly"] = {
            "reporting_month": "2026-07",
            "planned_percentage": None,
            "actual_percentage": None,
            "variance_percentage": None,
            "available": False,
            "message": "Monthly progress data is not available for July 2026.",
        }
        snap["physical_progress"]["monthly_available"] = False
        snap["physical_progress"]["history"] = [
            {"month": "2026-05", "planned_percentage": 45, "actual_percentage": 40}
        ]
        snap["physical_progress"]["cumulative"] = {
            "actual_percentage": 20.9,
            "scope_planned_quantity": 177.0,
            "scope_completed_quantity": 37.0,
            "scope_progress_percentage": 20.9,
            "available": True,
            "data_source": "scope",
        }
        out = validate_and_normalize_snapshot(snap)
        chart = out["physical_progress"]["chart"]
        self.assertEqual(chart["reporting_month"], "2026-07")
        self.assertEqual(chart["available_periods"], ["2026-05"])
        self.assertNotIn("2026-07", chart["available_periods"])
        self.assertNotIn("2026-07", chart["periods"])
        self.assertFalse(chart["reporting_month_available"])
        self.assertEqual(chart["render_mode"], "info_block")
        self.assertEqual(chart["planned"], [45])
        self.assertEqual(chart["actual"], [40])
        self.assertIn(
            "Monthly Construction Progress data is not available for July 2026",
            chart["message"] or "",
        )
        self.assertEqual(chart["cumulative_scope_progress"], 20.9)

        text = _pdf_text(render_pdf(out, meta={"version": 1}))
        self.assertIn("REPORTING PERIOD: JULY 2026", text)
        self.assertIn("May 2026", text)
        self.assertIn("45", text)
        self.assertIn("40", text)
        self.assertIn("not available", text.lower())
        self.assertIn("20.9", text)
        # Must not invent a July series point as 0
        self.assertNotIn("2026-07 — Planned 0", text)

    def test_may_june_july_cp_july_plotted(self):
        snap = _minimal_snapshot()
        out = validate_and_normalize_snapshot(snap)
        chart = out["physical_progress"]["chart"]
        self.assertEqual(chart["reporting_month"], "2026-07")
        self.assertTrue(chart["reporting_month_available"])
        self.assertEqual(chart["available_periods"][-1], "2026-07")
        self.assertEqual(chart["render_mode"], "chart")
        self.assertIn("2026-07", chart["periods"])

    def test_no_cp_records_info_block(self):
        snap = _minimal_snapshot()
        snap["physical_progress"]["monthly"] = {
            "reporting_month": "2026-07",
            "planned_percentage": None,
            "actual_percentage": None,
            "available": False,
            "message": "Monthly progress data is not available for July 2026.",
        }
        snap["physical_progress"]["monthly_available"] = False
        snap["physical_progress"]["history"] = []
        snap["physical_progress"]["cumulative"] = {
            "available": False,
            "actual_percentage": None,
        }
        out = validate_and_normalize_snapshot(snap)
        chart = out["physical_progress"]["chart"]
        self.assertEqual(chart["reporting_month"], "2026-07")
        self.assertEqual(chart["available_periods"], [])
        self.assertEqual(chart["render_mode"], "info_block")
        self.assertFalse(chart["reporting_month_available"])
        text = _pdf_text(render_pdf(out, meta={"version": 1}))
        self.assertIn("REPORTING PERIOD: JULY 2026", text)
        self.assertIn("No historical monthly Construction Progress", text)

    def test_multiple_historical_missing_july(self):
        snap = _minimal_snapshot()
        snap["physical_progress"]["monthly"] = {
            "reporting_month": "2026-07",
            "planned_percentage": None,
            "actual_percentage": None,
            "available": False,
            "message": "Monthly progress data is not available for July 2026.",
        }
        snap["physical_progress"]["monthly_available"] = False
        snap["physical_progress"]["history"] = [
            {"month": "2026-04", "planned_percentage": 20, "actual_percentage": 18},
            {"month": "2026-05", "planned_percentage": 30, "actual_percentage": 28},
            {"month": "2026-06", "planned_percentage": 35, "actual_percentage": 32},
        ]
        out = validate_and_normalize_snapshot(snap)
        chart = out["physical_progress"]["chart"]
        self.assertEqual(chart["reporting_month"], "2026-07")
        self.assertEqual(
            chart["available_periods"], ["2026-04", "2026-05", "2026-06"]
        )
        self.assertNotIn("2026-07", chart["available_periods"])
        self.assertFalse(chart["reporting_month_available"])
        self.assertEqual(chart["render_mode"], "chart")
        self.assertIn("July 2026", chart["message"] or "")
        text = _pdf_text(render_pdf(out, meta={"version": 1}))
        self.assertIn("REPORTING PERIOD: JULY 2026", text)
        self.assertIn("not available", text.lower())


class PhysicalChartPresentationDBTests(TestCase):
    def setUp(self):
        cache.clear()
        self.user = User.objects.create_user(
            username="mpr_chart_admin",
            email="mpr_chart@example.com",
            password="pass12345",
            is_staff=True,
            is_superuser=True,
        )
        self.project = Project.objects.create(
            name="B9300 Chart Gap Project",
            client_name="Test Client",
            location="Pune",
            status="active",
            project_start=date(2025, 1, 1),
            contract_finish=date(2026, 12, 31),
            team_lead=self.user,
        )
        self.period = parse_mpr_month("2026-07")
        self.cat = ScopeCategory.objects.create(name="Civil-Chart")
        self.sub = ScopeSubCategory.objects.create(
            category=self.cat, name="Excavation-Chart"
        )

    def _build(self):
        return MPRService(self.project, self.period).build()

    def test_scope_without_cp_keeps_cumulative_and_july_reporting_month(self):
        MonthlyScopeWork.objects.create(
            project=self.project,
            category=self.cat,
            subcategory=self.sub,
            month=date(2026, 5, 1),
            planned_quantity=Decimal("177"),
            cumulative_quantity=Decimal("37"),
            unit="cum",
        )
        ConstructionProgress.objects.create(
            projectName=self.project.name,
            progressMonth="2026-05",
            plannedProgress=45.0,
            actualProgress=40.0,
        )
        snap = self._build()
        chart = snap["physical_progress"]["chart"]
        cum = snap["physical_progress"]["cumulative"]
        self.assertEqual(chart["reporting_month"], "2026-07")
        self.assertEqual(chart["available_periods"], ["2026-05"])
        self.assertNotIn("2026-07", chart["available_periods"])
        self.assertFalse(chart["reporting_month_available"])
        self.assertEqual(chart["render_mode"], "info_block")
        self.assertAlmostEqual(cum["actual_percentage"], 20.9, places=1)
        self.assertEqual(chart["cumulative_scope_progress"], cum["actual_percentage"])

        text = _pdf_text(render_pdf(snap, meta={"version": 1}))
        self.assertIn("REPORTING PERIOD: JULY 2026", text)
        self.assertIn("May 2026", text)
        self.assertAlmostEqual(
            snap["executive_summary"]["physical_progress_billed_till_date"][
                "actual_percentage"
            ],
            20.9,
            places=1,
        )

        wb = load_workbook(BytesIO(render_excel(snap)))
        ws = wb["Physical Progress"]
        chart_rows = {
            (r[0], r[1]): r[2]
            for r in ws.iter_rows(min_row=1, values_only=True)
            if r and r[0] == "chart"
        }
        self.assertEqual(chart_rows.get(("chart", "reporting_month")), "2026-07")
        self.assertEqual(chart_rows.get(("chart", "available_periods")), "2026-05")
        self.assertFalse(chart_rows.get(("chart", "reporting_month_available")))
        self.assertEqual(chart_rows.get(("chart", "render_mode")), "info_block")
        self.assertNotIn("2026-07", str(chart_rows.get(("chart", "planned"))))
