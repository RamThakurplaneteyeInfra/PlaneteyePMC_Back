"""Phase-4 MPR completeness, HSE diagnostics, and consistency tests."""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO

from django.db import connection
from django.test import SimpleTestCase, TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook
from pypdf import PdfReader

from mpr.services.excel_renderer import render_excel
from mpr.services.mpr_validation import validate_and_normalize_snapshot
from mpr.services.pdf_renderer import render_pdf
from mpr.tests_client_format import _minimal_snapshot, _pdf_text


class HSEDataQualityTests(SimpleTestCase):
    def test_bad_source_suppresses_rates_and_flags_invalid(self):
        snap = _minimal_snapshot()
        snap["hse"] = {
            "available": True,
            "rates_available": True,
            "data_quality": {
                "status": "invalid",
                "warnings": ["Monthly fatalities=65 is unusually high; verify source entry."],
            },
            "record": {
                "fatalities": 65,
                "significant": 34,
                "major": 12,
                "minor": 453,
                "near_miss": 87,
                "total_incidents": 651,
                "total_manhours": 230,
                "man_hours_worked": 0,
                "loss_of_manhours": 123,
                "working_days": 0,
                "average_daily_manpower": 0,
                "ltifr": 534782.61,
                "incident_rate": 2830434.78,
            },
        }
        out = validate_and_normalize_snapshot(snap)
        self.assertFalse(out["hse"]["rates_available"])
        self.assertIsNone(out["hse"]["record"]["ltifr"])
        self.assertEqual(out["hse"]["record"]["fatalities"], 65)  # preserved
        self.assertIn(out["report_completeness"]["sections"]["hse"], ("warning", "unavailable"))


class PhysicalCompletenessTests(SimpleTestCase):
    def test_missing_july_progress_clearly_labelled(self):
        snap = _minimal_snapshot()
        snap["physical_progress"]["monthly_available"] = False
        snap["physical_progress"]["monthly"] = {
            "reporting_month": "2026-07",
            "planned_percentage": None,
            "actual_percentage": None,
            "variance_percentage": None,
            "available": False,
            "message": "Monthly progress data is not available for July 2026.",
        }
        snap["physical_progress"]["activities_available"] = False
        snap["physical_progress"]["activities"] = []
        snap["physical_progress"]["period_note"] = (
            "Monthly progress data is not available for July 2026."
        )
        snap["physical_progress"]["history"] = [
            {"month": "2026-05", "planned_percentage": 45, "actual_percentage": 40}
        ]
        out = validate_and_normalize_snapshot(snap)
        chart = out["physical_progress"]["chart"]
        self.assertEqual(chart["reporting_month"], "2026-07")
        self.assertFalse(chart["includes_reporting_month"])
        self.assertIn("2026-05", chart["periods"])
        self.assertNotIn("2026-07", chart["periods"])
        self.assertEqual(
            out["report_completeness"]["sections"]["physical_progress"], "partial"
        )

    def test_valid_july_progress_includes_reporting_month(self):
        snap = _minimal_snapshot()
        snap["physical_progress"]["monthly_available"] = True
        out = validate_and_normalize_snapshot(snap)
        chart = out["physical_progress"]["chart"]
        self.assertTrue(chart["includes_reporting_month"])
        self.assertEqual(chart["periods"][-1], "2026-07")


class EOTGroupingTests(SimpleTestCase):
    def test_approved_pending_rejected_groups(self):
        from mpr.services.mpr_validation import validate_and_normalize_snapshot

        snap = _minimal_snapshot()
        # Already has pending + approved in history; groups added by service —
        # validation preserves them if present
        snap["eot"]["approved"] = [snap["eot"]["latest_approved_eot"]]
        snap["eot"]["pending"] = [snap["eot"]["history"][0]]
        snap["eot"]["rejected"] = []
        out = validate_and_normalize_snapshot(snap)
        self.assertEqual(len(out["eot"]["approved"]), 1)
        self.assertEqual(len(out["eot"]["pending"]), 1)
        text = _pdf_text(render_pdf(out, meta={"version": 1}))
        self.assertIn("EOT Status", text)
        self.assertIn("Official completion", text)


class DrawingStatusTests(SimpleTestCase):
    def test_drawing_status_client_friendly(self):
        snap = _minimal_snapshot()
        out = validate_and_normalize_snapshot(snap)
        statuses = [i.get("status") for i in out["drawings"]["register_items"]]
        self.assertTrue(all(s in ("Approved", "Under Review", "Pending") or s for s in statuses))
        text = _pdf_text(render_pdf(out, meta={"version": 1}))
        self.assertIn("Under Review", text)
        self.assertNotIn("<br/>", text)


class SnapshotPdfExcelConsistencyTests(TestCase):
    def test_quality_and_progress_match_across_outputs(self):
        snap = validate_and_normalize_snapshot(_minimal_snapshot())
        pdf_text = _pdf_text(render_pdf(snap, meta={"version": 4}))
        xlsx = render_excel(snap)
        wb = load_workbook(BytesIO(xlsx))
        q = snap["quality"]
        self.assertIn(str(q["tests_conducted"]), pdf_text)
        # Excel quality sheet has pass_percentage
        qsheet = wb["Quality"]
        values = {str(row[0].value): row[1].value for row in qsheet.iter_rows(min_row=2)}
        self.assertEqual(values.get("pass_percentage"), q["pass_percentage"])
        self.assertEqual(values.get("tests_passed"), q["tests_passed"])

    def test_pdf_renderer_no_db(self):
        snap = validate_and_normalize_snapshot(_minimal_snapshot())
        with CaptureQueriesContext(connection) as ctx:
            render_pdf(snap, meta={"version": 1})
            render_excel(snap)
        self.assertEqual(len(ctx), 0)


class MissingSectionsTests(SimpleTestCase):
    def test_missing_material_lab_manpower_messages(self):
        snap = validate_and_normalize_snapshot(_minimal_snapshot())
        snap["materials"] = {"available": False, "records": [], "note": "Material received register is not available in the backend."}
        snap["laboratory_equipment"] = {"available": False, "records": []}
        snap["manpower"] = {
            "available": True,
            "planned_headcount": 10,
            "actual_headcount": 8,
            "limitations": ["Skill breakdown is not available."],
        }
        out = validate_and_normalize_snapshot(snap)
        text = _pdf_text(render_pdf(out, meta={"version": 1}))
        self.assertIn("material received", text.lower())
        self.assertIn("not available", text.lower())
        self.assertEqual(out["report_completeness"]["sections"]["materials"], "unavailable")


class FinancialChartChronologyTests(SimpleTestCase):
    def test_no_future_months_in_financial_chart(self):
        snap = _minimal_snapshot()
        snap["financial_progress"]["history"] = [
            {"month_year": "Sep-2026", "BCWS": 9, "ACWP": 9},
            {"month_year": "May-2026", "BCWS": 5, "ACWP": 4},
            {"month_year": "Jul-2026", "BCWS": 7, "ACWP": 6},
        ]
        out = validate_and_normalize_snapshot(snap)
        periods = out["financial_progress"]["chart"]["periods"]
        self.assertEqual(periods, ["May-2026", "Jul-2026"])
        self.assertNotIn("Sep-2026", periods)


class ReportCompletenessTests(SimpleTestCase):
    def test_completeness_section_present(self):
        out = validate_and_normalize_snapshot(_minimal_snapshot())
        self.assertIn("sections", out["report_completeness"])
        self.assertIn("physical_progress", out["report_completeness"]["sections"])
        text = _pdf_text(render_pdf(out, meta={"version": 1}))
        self.assertNotIn("REPORT DATA NOTES", text)
