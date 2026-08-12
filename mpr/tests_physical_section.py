"""Section 4.1 Physical Progress — monthly vs cumulative independence tests."""

from __future__ import annotations

from datetime import date
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext

from construction_progress.models.construction_progress import ConstructionProgress
from monthly_scope.models import MonthlyScopeWork, ScopeCategory, ScopeSubCategory
from mpr.services.excel_renderer import render_excel
from mpr.services.mpr_service import MPRService
from mpr.services.period import parse_mpr_month
from mpr.services.pdf_renderer import render_pdf
from projects.models import Project

User = get_user_model()


class PhysicalProgressSection41Tests(TestCase):
    def setUp(self):
        cache.clear()
        self.user = User.objects.create_user(
            username="mpr_phys_admin",
            email="mpr_phys@example.com",
            password="pass12345",
            is_staff=True,
            is_superuser=True,
        )
        self.project = Project.objects.create(
            name="B9200 Physical Progress Section Project",
            client_name="Test Client",
            location="Pune",
            status="active",
            project_start=date(2025, 1, 1),
            contract_finish=date(2026, 12, 31),
            team_lead=self.user,
        )
        self.month = "2026-07"
        self.period = parse_mpr_month(self.month)
        self.cat = ScopeCategory.objects.create(name="Civil-Phys")
        self.sub = ScopeSubCategory.objects.create(
            category=self.cat, name="Excavation-Phys"
        )

    def _build(self):
        return MPRService(self.project, self.period).build()

    def _scope(self, planned, cumulative, month=None):
        return MonthlyScopeWork.objects.create(
            project=self.project,
            category=self.cat,
            subcategory=self.sub,
            month=month or date(2026, 5, 1),
            planned_quantity=Decimal(str(planned)),
            cumulative_quantity=Decimal(str(cumulative)),
            unit="cum",
        )

    def _july_cp(self, planned=50, actual=40):
        return ConstructionProgress.objects.create(
            projectName=self.project.name,
            progressMonth=self.month,
            plannedProgress=float(planned),
            actualProgress=float(actual),
        )

    def test_july_cp_and_scope_both_available(self):
        self._july_cp(50, 40)
        self._scope(100, 40)
        snap = self._build()
        phys = snap["physical_progress"]
        monthly = phys["monthly"]
        cum = phys["cumulative"]
        self.assertTrue(monthly["available"])
        self.assertEqual(monthly["planned_percentage"], 50.0)
        self.assertEqual(monthly["actual_percentage"], 40.0)
        self.assertEqual(monthly["variance_percentage"], -10.0)
        self.assertTrue(cum["available"])
        self.assertEqual(cum["actual_percentage"], 40.0)
        self.assertEqual(cum["scope_planned_quantity"], 100.0)
        self.assertEqual(cum["scope_completed_quantity"], 40.0)
        self.assertEqual(cum["scope_progress_percentage"], 40.0)
        self.assertIsNone(monthly.get("message"))
        self.assertNotIn("Scope activity data is not available", phys.get("activities_note") or "")

    def test_july_cp_missing_scope_exists_satis_like(self):
        ConstructionProgress.objects.create(
            projectName=self.project.name,
            progressMonth="2026-05",
            plannedProgress=45.0,
            actualProgress=40.0,
        )
        self._scope(177, 37)
        snap = self._build()
        phys = snap["physical_progress"]
        monthly = phys["monthly"]
        cum = phys["cumulative"]

        self.assertFalse(monthly["available"])
        self.assertIsNone(monthly["planned_percentage"])
        self.assertIsNone(monthly["actual_percentage"])
        self.assertIsNone(monthly["variance_percentage"])
        self.assertIn("Monthly progress data is not available for July 2026", monthly["message"])
        self.assertNotIn("Scope activity data is not available", monthly["message"])
        self.assertNotIn(
            "Scope activity data is not available",
            phys.get("activities_note") or "",
        )
        self.assertNotIn(
            "Scope activity data is not available",
            phys.get("period_note") or "",
        )

        self.assertTrue(cum["available"])
        self.assertAlmostEqual(cum["actual_percentage"], 20.9, places=1)
        self.assertEqual(cum["scope_planned_quantity"], 177.0)
        self.assertEqual(cum["scope_completed_quantity"], 37.0)
        self.assertAlmostEqual(cum["scope_progress_percentage"], 20.9, places=1)

        # May must not become July
        self.assertNotEqual(monthly["planned_percentage"], 45.0)
        self.assertNotEqual(monthly["actual_percentage"], 40.0)
        chart = phys["chart"]
        self.assertFalse(chart["includes_reporting_month"])
        self.assertIn("2026-05", chart["periods"])
        self.assertNotIn("2026-07", chart["periods"])

        # Executive summary uses same cumulative actual
        billed = snap["executive_summary"]["physical_progress_billed_till_date"]
        self.assertAlmostEqual(
            billed["actual_percentage"], cum["actual_percentage"], places=2
        )

        from io import BytesIO

        from pypdf import PdfReader

        pdf = render_pdf(snap, meta={"version": 1})
        body = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(pdf)).pages)
        self.assertIn("Monthly progress data is not available for July 2026", body)
        self.assertNotIn("Scope activity data is not available for 2026-07", body)
        self.assertTrue("20.9" in body)

    def test_no_cp_no_scope(self):
        snap = self._build()
        phys = snap["physical_progress"]
        self.assertFalse(phys["monthly"]["available"])
        self.assertFalse(phys["cumulative"]["available"])
        self.assertIsNone(phys["cumulative"]["actual_percentage"])

    def test_cp_exists_no_scope(self):
        self._july_cp(55, 48)
        snap = self._build()
        phys = snap["physical_progress"]
        self.assertTrue(phys["monthly"]["available"])
        self.assertEqual(phys["monthly"]["planned_percentage"], 55.0)
        self.assertEqual(phys["monthly"]["actual_percentage"], 48.0)
        self.assertFalse(phys["cumulative"]["available"])
        self.assertIsNone(phys["cumulative"]["actual_percentage"])

    def test_scope_planned_zero_no_division(self):
        self._scope(0, 0)
        snap = self._build()
        cum = snap["physical_progress"]["cumulative"]
        self.assertTrue(cum["available"])
        self.assertEqual(cum["scope_planned_quantity"], 0.0)
        self.assertEqual(cum["scope_completed_quantity"], 0.0)
        self.assertIsNone(cum["scope_progress_percentage"])
        self.assertIsNone(cum["actual_percentage"])

    def test_completed_exceeds_planned_clamps_pct_keeps_quantities(self):
        self._scope(100, 150)
        snap = self._build()
        cum = snap["physical_progress"]["cumulative"]
        self.assertEqual(cum["scope_planned_quantity"], 100.0)
        self.assertEqual(cum["scope_completed_quantity"], 150.0)
        # Overview rule clamps percentage at 100
        self.assertEqual(cum["actual_percentage"], 100.0)
        self.assertEqual(cum["scope_progress_percentage"], 100.0)

    def test_excel_matches_snapshot_cumulative(self):
        self._scope(177, 37)
        snap = self._build()
        cum = snap["physical_progress"]["cumulative"]
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(render_excel(snap)))
        ws = wb["Physical Progress"]
        found_summary = False
        found_block = False
        for row in ws.iter_rows(min_row=1, values_only=True):
            if not row:
                continue
            if row[0] == "cumulative_actual_percentage":
                self.assertAlmostEqual(
                    float(row[1]), float(cum["actual_percentage"]), places=1
                )
                found_summary = True
            if row[0] == "cumulative" and row[1] == "actual_percentage":
                self.assertAlmostEqual(
                    float(row[2]), float(cum["actual_percentage"]), places=1
                )
                found_block = True
        self.assertTrue(found_summary)
        self.assertTrue(found_block)

    def test_no_extra_n_plus_one_on_build(self):
        self._july_cp()
        self._scope(100, 40)
        # Warm / establish baseline
        MPRService(self.project, self.period).build()
        with CaptureQueriesContext(connection) as ctx:
            MPRService(self.project, self.period).build()
        # Absolute budget — must stay bounded (no per-scope-row queries)
        self.assertLess(len(ctx), 80)
