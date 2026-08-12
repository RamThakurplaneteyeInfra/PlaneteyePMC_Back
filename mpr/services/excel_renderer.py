"""
Excel workbook renderer for MPR — snapshot_json only (no DB queries).
"""

from __future__ import annotations

from io import BytesIO
from typing import Any


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return str(value)
    return value


def _write_sheet(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([_cell(c) for c in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, _ in enumerate(headers, start=1):
        letter = ws.cell(1, idx).column_letter
        ws.column_dimensions[letter].width = min(36, max(12, len(headers[idx - 1]) + 4))


def _kv_rows(mapping: dict | None) -> list[list[Any]]:
    if not mapping:
        return []
    rows = []
    for key, value in mapping.items():
        if isinstance(value, dict):
            for k2, v2 in value.items():
                rows.append([f"{key}.{k2}", v2])
        elif isinstance(value, list):
            rows.append([key, f"{len(value)} items"])
        else:
            rows.append([key, value])
    return rows


def render_excel(snapshot: dict) -> bytes:
    """Return .xlsx bytes from an MPR snapshot dict."""
    from openpyxl import Workbook

    wb = Workbook()

    # 1. Project Summary
    ws = wb.active
    ws.title = "Project Summary"
    project = snapshot.get("project") or {}
    period = snapshot.get("reporting_period") or {}
    _write_sheet(
        ws,
        ["Field", "Value"],
        _kv_rows({**project, **{f"period_{k}": v for k, v in period.items()}}),
    )

    # 2. Key Indicators
    ws = wb.create_sheet("Key Indicators")
    _write_sheet(ws, ["Indicator", "Value"], _kv_rows(snapshot.get("key_indicators") or {}))

    # 3. Physical Progress
    ws = wb.create_sheet("Physical Progress")
    phys = snapshot.get("physical_progress") or {}
    monthly = phys.get("monthly") or {}
    cumulative = phys.get("cumulative") or {}
    scope = phys.get("scope_progress") or {}
    rows = [
        ["reporting_month", phys.get("reporting_month"), ""],
        ["reporting_month_label", phys.get("reporting_month_label"), ""],
        ["monthly_available", monthly.get("available", phys.get("monthly_available")), ""],
        ["monthly_message", monthly.get("message") or phys.get("period_note"), ""],
        ["cumulative_available", cumulative.get("available"), ""],
        ["cumulative_actual_percentage", cumulative.get("actual_percentage"), ""],
        ["scope_planned_quantity", cumulative.get("scope_planned_quantity") or scope.get("planned_quantity"), ""],
        ["scope_completed_quantity", cumulative.get("scope_completed_quantity") or scope.get("cumulative_quantity"), ""],
        ["scope_progress_percentage", cumulative.get("scope_progress_percentage") or scope.get("progress_percentage"), ""],
        ["activities_available", phys.get("activities_available"), ""],
    ]
    chart = phys.get("chart") or {}
    rows.extend(
        [
            ["chart", "reporting_month", chart.get("reporting_month")],
            ["chart", "reporting_month_label", chart.get("reporting_month_label")],
            ["chart", "available_periods", ",".join(chart.get("available_periods") or chart.get("periods") or [])],
            ["chart", "reporting_month_available", chart.get("reporting_month_available")],
            ["chart", "render_mode", chart.get("render_mode")],
            ["chart", "message", chart.get("message")],
            ["chart", "planned", str(chart.get("planned") or [])],
            ["chart", "actual", str(chart.get("actual") or [])],
            ["chart", "cumulative_scope_progress", chart.get("cumulative_scope_progress")],
        ]
    )
    for section in ("monthly", "cumulative", "scope_progress"):
        block = phys.get(section) or {}
        for k, v in block.items():
            rows.append([section, k, v])
    _write_sheet(ws, ["Section", "Field", "Value"], rows)
    ws_act = wb.create_sheet("Physical Activities")
    _write_sheet(
        ws_act,
        ["Sr", "Item", "Unit", "Total", "Completed", "Balance", "%", "Remarks"],
        [
            [
                a.get("sr_no"),
                a.get("item"),
                a.get("unit"),
                a.get("total"),
                a.get("completed"),
                a.get("balance"),
                a.get("percent_achieved"),
                a.get("remarks"),
            ]
            for a in (phys.get("activities") or [])
        ],
    )

    # 3b. Scope Items detail
    ws_scope = wb.create_sheet("Scope Items")
    phys = snapshot.get("physical_progress") or {}
    _write_sheet(
        ws_scope,
        [
            "Sr",
            "Category",
            "Item",
            "Unit",
            "Planned",
            "Completed",
            "Balance",
            "Progress %",
            "Status",
            "Scope Month",
        ],
        [
            [
                a.get("sr_no"),
                a.get("category"),
                a.get("item"),
                a.get("unit"),
                a.get("total"),
                a.get("completed"),
                a.get("balance"),
                a.get("percent_achieved"),
                a.get("status"),
                a.get("scope_month"),
            ]
            for a in (phys.get("scope_items") or [])
        ],
    )
    ws_cat = wb.create_sheet("Scope Categories")
    _write_sheet(
        ws_cat,
        ["Category", "Items", "Planned", "Completed", "Progress %"],
        [
            [
                c.get("category"),
                c.get("item_count"),
                c.get("planned_quantity"),
                c.get("completed_quantity"),
                c.get("progress_percentage"),
            ]
            for c in (phys.get("scope_category_summary") or [])
        ],
    )

    # 4. EOT
    ws = wb.create_sheet("EOT")
    eot = snapshot.get("eot") or {}
    history = eot.get("history") or []
    _write_sheet(
        ws,
        [
            "eot_number",
            "extension_days",
            "reason",
            "approval_date",
            "original_completion_date",
            "revised_completion_date",
            "status",
        ],
        [
            [
                h.get("eot_number"),
                h.get("extension_days"),
                h.get("reason"),
                h.get("approval_date"),
                h.get("original_completion_date"),
                h.get("revised_completion_date"),
                h.get("status"),
            ]
            for h in history
        ],
    )

    # 5. Financial
    ws = wb.create_sheet("Financial")
    fin = snapshot.get("financial_progress") or {}
    rows = []
    for section in ("contract", "evm", "cashflow"):
        block = fin.get(section) or {}
        if isinstance(block, dict):
            for k, v in block.items():
                rows.append([section, k, v])
    for inv in (fin.get("invoicing") or {}).get("records") or []:
        for k, v in inv.items():
            rows.append(["invoicing", k, v])
    _write_sheet(ws, ["Section", "Field", "Value"], rows)

    # 6. Correspondence
    ws = wb.create_sheet("Correspondence")
    corr = snapshot.get("correspondence") or {}
    summary = corr.get("summary") or {}
    _write_sheet(ws, ["Metric", "Value"], [[k, v] for k, v in summary.items()])
    ws2_rows = [
        [
            r.get("sr_no"),
            r.get("received_date"),
            (r.get("description") or "")[:120],
            r.get("delivered_status"),
        ]
        for r in (corr.get("important_records") or [])
    ]
    ws_docs = wb.create_sheet("Correspondence Docs")
    _write_sheet(ws_docs, ["Letter No", "Date", "Subject", "Status"], ws2_rows)

    # 7. Drawings
    ws = wb.create_sheet("Drawings")
    drawings = snapshot.get("drawings") or {}
    _write_sheet(
        ws,
        ["Metric", "Value"],
        [
            ["total_register", drawings.get("total_register")],
            ["submitted", drawings.get("submitted")],
            ["approved", drawings.get("approved")],
            ["pending", drawings.get("pending")],
            ["approval_rate", drawings.get("approval_rate")],
        ],
    )
    ws_d = wb.create_sheet("Drawings Pending")
    _write_sheet(
        ws_d,
        ["Sr", "Title", "Revision", "Submitted", "Remarks"],
        [
            [
                d.get("sr_no"),
                d.get("drawing_name"),
                d.get("revision"),
                d.get("submitted_date"),
                d.get("remarks"),
            ]
            for d in (drawings.get("pending_items") or [])
        ],
    )

    # 8. Bottlenecks
    ws = wb.create_sheet("Bottlenecks")
    bn = snapshot.get("bottlenecks") or {}
    _write_sheet(
        ws,
        [
            "Issue",
            "Category",
            "Responsible",
            "Date Raised",
            "Target",
            "Status",
            "Days Open",
            "Action",
        ],
        [
            [
                (r.get("description") or "")[:120],
                r.get("category") or r.get("type"),
                r.get("assigned_to"),
                r.get("date_raised") or r.get("created_at"),
                r.get("target_date"),
                r.get("status"),
                r.get("days_open") if r.get("days_open") is not None else r.get("ageing_days"),
                r.get("action_required"),
            ]
            for r in (bn.get("records") or [])
        ],
    )

    # 9. Quality
    ws = wb.create_sheet("Quality")
    _write_sheet(ws, ["Field", "Value"], _kv_rows(snapshot.get("quality") or {}))

    # 10. HSE
    ws = wb.create_sheet("HSE")
    hse = snapshot.get("hse") or {}
    record = hse.get("record") if isinstance(hse.get("record"), dict) else {}
    hse_rows = _kv_rows(record if isinstance(record, dict) else {})
    hse_rows.append(["rates_available", hse.get("rates_available")])
    dq = hse.get("data_quality") or {}
    hse_rows.append(["data_quality.status", dq.get("status")])
    hse_rows.append(["data_quality.warnings", "; ".join(dq.get("warnings") or [])])
    _write_sheet(ws, ["Field", "Value"], hse_rows)

    # 11. Manpower
    ws = wb.create_sheet("Manpower")
    _write_sheet(ws, ["Field", "Value"], _kv_rows(snapshot.get("manpower") or {}))

    # 12. Equipment
    ws = wb.create_sheet("Equipment")
    eq = snapshot.get("equipment") or {}
    kpi = eq.get("kpi") or eq.get("deployment_summary") or {}
    summary_rows = _kv_rows(kpi if isinstance(kpi, dict) else {})
    summary_rows.extend(
        [
            ["count", eq.get("count")],
            ["total_quantity", eq.get("total_quantity")],
            ["source_report_date", eq.get("source_report_date")],
        ]
    )
    _write_sheet(ws, ["Field", "Value"], summary_rows)
    ws_inv = wb.create_sheet("Plant Machinery")
    _write_sheet(
        ws_inv,
        ["Sr", "Name", "Unit", "Qty", "Status", "Remark", "Category"],
        [
            [
                item.get("sr_no"),
                item.get("name"),
                item.get("unit"),
                item.get("qty"),
                item.get("status"),
                item.get("remark"),
                item.get("category"),
            ]
            for item in (eq.get("inventory") or [])
        ],
    )

    # 13. Site Photos
    ws = wb.create_sheet("Site Photos")
    photos = (snapshot.get("site_photos") or {}).get("photos") or []
    _write_sheet(
        ws,
        ["Title", "URL", "Created", "Month", "Year"],
        [
            [
                p.get("title"),
                p.get("image_url"),
                p.get("created_at"),
                p.get("month"),
                p.get("year"),
            ]
            for p in photos
        ],
    )

    # 14. Data Availability / Completeness
    ws = wb.create_sheet("Data Availability")
    avail = snapshot.get("data_availability") or {}
    rows = []
    for section, info in avail.items():
        if isinstance(info, dict):
            rows.append(
                [
                    section,
                    info.get("status"),
                    ", ".join(info.get("missing") or []),
                    info.get("note") or "",
                ]
            )
        else:
            rows.append([section, info, "", ""])
    _write_sheet(ws, ["Section", "Status", "Missing", "Note"], rows)
    ws_c = wb.create_sheet("Report Completeness")
    comp = snapshot.get("report_completeness") or {}
    _write_sheet(
        ws_c,
        ["Section", "Status"],
        [[k, v] for k, v in (comp.get("sections") or {}).items()],
    )

    # BG sheet (bonus)
    ws = wb.create_sheet("BG")
    bg = snapshot.get("bg") or {}
    _write_sheet(
        ws,
        ["Name", "Type", "Due", "Updated", "Status", "Remarks"],
        [
            [
                r.get("bg_name"),
                r.get("bg_type"),
                r.get("due_date"),
                r.get("updated_date"),
                r.get("status"),
                r.get("remarks"),
            ]
            for r in (bg.get("records") or [])
        ],
    )

    # Meetings / Next month
    ws = wb.create_sheet("Meetings")
    _write_sheet(
        ws,
        ["Date", "Type", "Title", "Description"],
        [
            [
                m.get("meeting_date"),
                m.get("meeting_type"),
                m.get("title"),
                (m.get("description") or "")[:160],
            ]
            for m in ((snapshot.get("meetings") or {}).get("records") or [])
        ],
    )
    ws = wb.create_sheet("Next Month Programme")
    _write_sheet(
        ws,
        ["Sr", "Activity", "Target", "Unit", "Planned Completion", "Remarks"],
        [
            [
                r.get("sr_no"),
                r.get("activity"),
                r.get("target"),
                r.get("unit"),
                r.get("planned_completion"),
                r.get("remarks"),
            ]
            for r in ((snapshot.get("next_month_program") or {}).get("records") or [])
        ],
    )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
